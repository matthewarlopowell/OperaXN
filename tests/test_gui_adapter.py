"""operaxn.input adapter: the GUI contract over the core pipeline."""
import logging
import os
import shutil
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest

pytestmark = pytest.mark.gui
pytest.importorskip("tkinter")

import h5py

import core
import operaxn
from operaxn import input as gui_input
from operaxn.dialog import UploadOptions, UploadOptionsDialog
from operaxn.gui import OPERAXN
from operaxn.main import main as cli_main
from operaxn.input import (
    DataSourceType,
    TimeMethod,
    add_standard_echem_files,
    cleanup_session_cache,
    experiment_metadata_rows,
    export_nxs,
    get_correlated_data,
    get_experiment_metadata,
    get_experiment_stem,
    get_loaded_data_source,
    get_loaded_nxs_path,
    get_standard_echem,
    make_echem_arrays,
    make_neutron_arrays,
    make_oned_arrays,
    make_twod_arrays,
    process_paths,
    reset_session,
)

import _builders
from _helpers import needs_fabio


def _reset_session():
    """Restore operaxn.input._session to its pristine module state."""
    gui_input._session.update({"nxs_path": None, "input_paths": None,
                               "exported": False, "standard_echem": [],
                               "stem": None, "data_source": None})
    gui_input._session.pop("global_metadata", None)


@pytest.fixture(autouse=True)
def _own_session():
    """Give each test a fresh module-global session and drop its cache .nxs after."""
    _reset_session()
    yield
    if gui_input._session.get("input_paths") is not None:
        gui_input._session["exported"] = False
    cleanup_session_cache()
    _reset_session()


def _discard(path):
    """Remove a cache file the test left behind in the shared cache dir."""
    if path and os.path.isfile(path):
        os.remove(path)


def _write_edf(path):
    """Write an 8x8 EDF image whose Date pairs it with in-house scan 1."""
    import fabio
    img = fabio.edfimage.EdfImage(
        data=np.arange(64, dtype=np.float32).reshape(8, 8) + 1.0,
        header={"Date": _builders.SCAN_TIMES[0], "ExposureTime": "120.0",
                "WaveLength": "1.54"})
    img.write(path)


@pytest.fixture
def inhouse_edf_dir(inhouse_dir):
    """Return the in-house dataset plus an EDF 2D image paired with scan 1 by Date."""
    _write_edf(os.path.join(inhouse_dir, "image_001.edf"))
    return inhouse_dir


def _std_echem_file(path, **overrides):
    """Write a standard (non-operando) echem file distinct from the operando series."""
    options = dict(n_rows=10, v0=4.0, dv=0.01, i0=0.5, di=0.0,
                   date="06/02/2024", hour=9)
    options.update(overrides)
    _builders.write_echem_txt(path, **options)


def test_package_exports_core_enums():
    """Full package import succeeds and config re-exports the core enum."""
    assert operaxn.config.DataSourceType is core.DataSourceType


def test_raw_upload_absolute(inhouse_dir):
    """Raw in-house upload yields correlated scan dicts with midpoint matching."""
    progress = []
    scans, _, method = process_paths(
        [inhouse_dir], progress.append,
        time_method=TimeMethod.ABSOLUTE, data_source=DataSourceType.INHOUSE)
    assert method == "absolute", method
    assert len(scans) == 3, f"got {len(scans)}"
    assert len(progress) > 0
    s1 = scans[0]
    assert s1["echem"] == 3.76, f"got {s1['echem']}"
    assert s1["echem_timestamp"] == "2024-02-05 10:06:00", \
        f"got {s1['echem_timestamp']}"
    assert s1["oned"] == "scan_001.dat", f"got {s1['oned']}"
    assert s1["timestamp_for_correlation"] == pd.Timestamp("2024-02-05 10:06:00"), \
        f"got {s1['timestamp_for_correlation']}"


@needs_fabio
def test_raw_upload_pairs_edf_and_lazy_loads(inhouse_edf_dir):
    """Scan 1 gains a 2D reference and make_twod_arrays lazy-loads the EDF."""
    scans, _, _ = process_paths(
        [inhouse_edf_dir], time_method=TimeMethod.ABSOLUTE,
        data_source=DataSourceType.INHOUSE)
    assert bool(scans[0]["twod"]), f"got {scans[0]['twod']}"
    twod_arrays = make_twod_arrays(scans)
    assert 1 in twod_arrays and "image" in twod_arrays[1], \
        str(twod_arrays.get(1, {}).keys())
    assert twod_arrays[1]["image"].shape == (8, 8)


def test_make_oned_arrays(inhouse_dir):
    """1D plot arrays carry exact data and the correlated voltage per scan."""
    scans, _, _ = process_paths(
        [inhouse_dir], time_method=TimeMethod.ABSOLUTE,
        data_source=DataSourceType.INHOUSE)
    oned_arrays = make_oned_arrays(scans)
    assert set(oned_arrays.keys()) == {1, 2, 3}
    assert np.allclose(oned_arrays[1]["x"], _builders.ONED_X)
    assert np.allclose(oned_arrays[1]["y"], _builders.inhouse_scan_y(1))
    assert oned_arrays[1]["echem"] == 3.76


def test_make_echem_arrays_absolute(inhouse_dir):
    """Echem arrays cover all points with x as seconds from the first point."""
    _, echem_df, method = process_paths(
        [inhouse_dir], time_method=TimeMethod.ABSOLUTE,
        data_source=DataSourceType.INHOUSE)
    echem_arrays = make_echem_arrays(echem_df, method)
    assert len(echem_arrays["x"]) == 40
    assert echem_arrays["x"][0] == 0.0


def test_get_correlated_data(inhouse_dir):
    """get_correlated_data bundles the 1D arrays and voltage for one scan."""
    scans, echem_df, _ = process_paths(
        [inhouse_dir], time_method=TimeMethod.ABSOLUTE,
        data_source=DataSourceType.INHOUSE)
    corr = get_correlated_data(scans, echem_df, 1)
    assert corr["oned"] is not None
    assert np.allclose(corr["oned"]["x"], _builders.ONED_X)
    assert corr["echem_value"] == 3.76


def test_direct_nxs_upload_matches_raw(inhouse_dir, tmp_path):
    """Opening the generated .nxs directly reproduces the raw-upload results."""
    scans, echem_df, _ = process_paths(
        [inhouse_dir], time_method=TimeMethod.ABSOLUTE,
        data_source=DataSourceType.INHOUSE)
    oned_arrays = make_oned_arrays(scans)
    nxs_copy = str(tmp_path / "shared_experiment.nxs")
    shutil.copy(core.default_cache_path([inhouse_dir]), nxs_copy)

    scans2, echem_df2, _ = process_paths(
        [nxs_copy], time_method=TimeMethod.ABSOLUTE,
        data_source=DataSourceType.INHOUSE)
    assert len(scans2) == 3, f"got {len(scans2)}"
    assert [s["echem"] for s in scans2] == [s["echem"] for s in scans]
    assert np.allclose(make_oned_arrays(scans2)[1]["y"], oned_arrays[1]["y"])
    assert np.allclose(echem_df2["echem_data"].values,
                       echem_df["echem_data"].values)


def test_relative_view_transform(inhouse_dir):
    """Relative view rewrites timestamps as offsets and re-zeroes matching."""
    scans_r, echem_r, method_r = process_paths(
        [inhouse_dir], time_method=TimeMethod.RELATIVE,
        data_source=DataSourceType.INHOUSE)
    assert method_r == "relative", method_r
    assert scans_r[0]["timestamp"] == "00:00:00", f"got {scans_r[0]['timestamp']}"
    assert scans_r[1]["timestamp"] == "00:05:00", f"got {scans_r[1]['timestamp']}"
    assert scans_r[0]["original_timestamp"] == "2024-02-05 10:05:00", \
        f"got {scans_r[0]['original_timestamp']}"
    assert echem_r["timestamp"].iloc[0] == "00:00:00", \
        f"got {echem_r['timestamp'].iloc[0]}"
    assert scans_r[0]["echem"] == 3.70, f"got {scans_r[0]['echem']}"

    rel_echem_arrays = make_echem_arrays(echem_r, method_r)
    assert rel_echem_arrays["x"][5] == 300.0, f"got {rel_echem_arrays['x'][5]}"


def test_neutron_adapter(neutron_dir):
    """Neutron upload yields one correlated scan with per-bank display names."""
    scans_n, _, _ = process_paths(
        [neutron_dir], time_method=TimeMethod.ABSOLUTE,
        data_source=DataSourceType.NEUTRON)
    assert len(scans_n) == 1, f"got {len(scans_n)}"
    sn = scans_n[0]
    assert sn["echem"] == 3.85, f"got {sn['echem']}"
    assert sn["neutron_files"] is not None, str(sn["neutron_files"])
    assert sn["neutron_files"]["1"]["tof"] == "123456-1-0.dat", \
        str(sn["neutron_files"])
    assert sn["neutron_files"]["1"]["d"] == "123456-1-d-0.dat", \
        str(sn["neutron_files"])

    n_arrays = make_neutron_arrays(scans_n)
    assert set(n_arrays[1].keys()) == {"1", "2"}
    assert np.allclose(n_arrays[1]["1"]["tof"]["x"], _builders.TOF_X)


@needs_fabio
def test_embedded_twod_survives_export_and_source_deletion(inhouse_edf_dir, tmp_path):
    """Images embedded at upload survive the copy export and load once the source EDF is gone."""
    process_paths([inhouse_edf_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE, include_2d_images=True)
    exported = str(tmp_path / "exported_embedded.nxs")
    ok, msgs = export_nxs(exported)
    assert ok, str(msgs)
    os.unlink(os.path.join(inhouse_edf_dir, "image_001.edf"))

    model = core.load(exported)
    assert model.scans[0].twod_embedded and model.scans[0].twod.shape == (8, 8)
    scans, _, _ = process_paths([exported], time_method=TimeMethod.ABSOLUTE,
                                data_source=DataSourceType.INHOUSE)
    twod = make_twod_arrays(scans)
    assert 1 in twod and "image" in twod[1], str(twod.get(1, {}).keys())
    assert twod[1]["image"].shape == (8, 8)
    _discard(core.default_cache_path([inhouse_edf_dir]))


def test_export_copies_canonical_file(inhouse_dir, tmp_path):
    """Export copies the generated file with scans and correlated voltages."""
    process_paths([inhouse_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    export_dest = str(tmp_path / "exported.nxs")
    ok, msgs = export_nxs(export_dest)
    assert ok, str(msgs)
    exported_model = core.load(export_dest)
    assert len(exported_model.scans) == 3
    assert exported_model.scans[0].echem == 3.76


@needs_fabio
def test_export_default_keeps_twod_by_reference(inhouse_edf_dir, tmp_path):
    """Default generation stores 2D by reference, and export preserves that."""
    process_paths([inhouse_edf_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    export_dest = str(tmp_path / "exported.nxs")
    ok, msgs = export_nxs(export_dest)
    assert ok, str(msgs)
    exported_model = core.load(export_dest)
    assert not exported_model.scans[0].twod_embedded
    assert exported_model.scans[0].twod is None


def test_direct_nxs_load_survives_cleanup_and_exports(inhouse_dir, tmp_path):
    """A directly opened user .nxs is never deleted by session cleanup and still exports."""
    process_paths([inhouse_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    kept_export = str(tmp_path / "kept.nxs")
    export_nxs(kept_export)

    process_paths([kept_export], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    cleanup_session_cache()
    assert os.path.isfile(kept_export)
    copy2 = str(tmp_path / "copy_of_direct.nxs")
    ok, _ = export_nxs(copy2)
    assert ok and os.path.isfile(copy2)
    _discard(core.default_cache_path([inhouse_dir]))


def test_cache_unsaved_removed_on_exit(inhouse_dir):
    """The unsaved session cache .nxs is deleted by cleanup on app exit."""
    process_paths([inhouse_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    cache_file = get_loaded_nxs_path()
    assert cache_file and os.path.isfile(cache_file)
    cleanup_session_cache()
    assert not os.path.exists(cache_file)


@pytest.mark.display
def test_closing_main_window_cleans_session_cache(tk_root, inhouse_dir):
    """The main window's registered close handler deletes the unsaved session cache."""
    from tkinter import Toplevel
    process_paths([inhouse_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    cache_file = get_loaded_nxs_path()
    assert cache_file and os.path.isfile(cache_file)
    win = Toplevel(tk_root)  # the handler destroys its master; keep the shared root
    OPERAXN(master=win)
    win.tk.eval(win.protocol("WM_DELETE_WINDOW"))
    assert not os.path.exists(cache_file), "exit cleanup must run on window close"
    assert not win.winfo_exists()


@pytest.mark.display
def test_declining_to_clear_cancels_the_upload(tk_root, monkeypatch):
    """Answering No to "Clear all data?" keeps the loaded scans and starts no load."""
    from tkinter import Toplevel, messagebox
    import operaxn.gui as gui_module
    win = Toplevel(tk_root)
    try:
        app = OPERAXN(master=win)
        loaded = [{"scan_num": 1}]
        app.state.scans = loaded
        options = UploadOptions(paths=["/new/data"], data_source=DataSourceType.INHOUSE,
                                time_method=TimeMethod.ABSOLUTE, display_size=0,
                                include_2d=False)
        monkeypatch.setattr(gui_module, "UploadOptionsDialog",
                            lambda *_a, **_k: SimpleNamespace(get_result=lambda: options))
        started = []
        monkeypatch.setattr(app, "_process_files_async", started.append)
        monkeypatch.setattr(messagebox, "askyesno", lambda *_a, **_k: False)
        app._upload_files()
        assert started == [] and app.state.scans is loaded, "No must cancel the load"
        monkeypatch.setattr(messagebox, "askyesno", lambda *_a, **_k: True)
        app._upload_files()
        assert started == [options] and not app.state.scans
    finally:
        if win.winfo_exists():
            win.destroy()


def test_cache_exported_file_kept_on_exit(inhouse_dir, tmp_path):
    """An exported session cache .nxs survives cleanup on app exit."""
    process_paths([inhouse_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    cache_file = get_loaded_nxs_path()
    export_nxs(str(tmp_path / "kept.nxs"))
    cleanup_session_cache()
    assert os.path.isfile(cache_file)


# ============================================================================
# ZIP 2D sidecar lifecycle (persist -> resolve -> export -> cleanup)
# ============================================================================

def _write_twod_zip(tmp_path):
    """Write a ZIP of one XRD .dat and one EDF image at one timestamp and return its path."""
    return _builders.write_twod_zip(str(tmp_path / "cell.zip"), tmp_path)


@needs_fabio
def test_cache_zip_images_dir_cleaned(tmp_path):
    """The images dir persisted for a ZIP input is removed with its cache file."""
    zpath = _write_twod_zip(tmp_path)
    process_paths([zpath], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    cache_file = get_loaded_nxs_path()
    images_dir = os.path.splitext(cache_file)[0] + "_images"
    assert os.path.isdir(images_dir), "images dir not persisted for ZIP input"
    cleanup_session_cache()
    assert not os.path.exists(cache_file)
    assert not os.path.isdir(images_dir)


@needs_fabio
def test_exported_zip_twod_resolves_after_cache_cleanup(tmp_path):
    """Export copies the ZIP 2D sidecar, and 2D resolves from it once the cache copy is gone."""
    zpath = _write_twod_zip(tmp_path)
    process_paths([zpath], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    cache_images = os.path.splitext(get_loaded_nxs_path())[0] + "_images"
    dest = str(tmp_path / "kept.nxs")
    ok, _ = export_nxs(dest)
    assert ok
    edfs = [f for _, _, fs in os.walk(str(tmp_path / "kept_images"))
            for f in fs if f.endswith(".edf")]
    assert edfs == ["image_001.edf"], f"sidecar not copied on export: {edfs}"
    shutil.rmtree(cache_images)  # simulate the cache being cleaned later
    scans, _, _ = process_paths([dest], time_method=TimeMethod.ABSOLUTE,
                                data_source=DataSourceType.INHOUSE)
    arrays = make_twod_arrays(scans)
    assert 1 in arrays and arrays[1]["image"].shape == (8, 8), \
        "2D not resolved from the exported sibling images dir"


@needs_fabio
def test_exported_session_cleanup_drops_cache_sidecar_keeps_nxs(tmp_path):
    """Exported-session cleanup keeps the cache .nxs but drops the sidecar the export has copied."""
    zpath = _write_twod_zip(tmp_path)
    process_paths([zpath], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    cache_file = get_loaded_nxs_path()
    cache_images = os.path.splitext(cache_file)[0] + "_images"
    dest = str(tmp_path / "kept.nxs")
    ok, _ = export_nxs(dest)
    assert ok and os.path.isdir(cache_images)
    cleanup_session_cache()  # exported flag still True: app-exit behaviour
    assert os.path.isfile(cache_file), "exported cache .nxs must be kept"
    assert not os.path.isdir(cache_images), "cache sidecar must not leak"
    assert os.path.isdir(str(tmp_path / "kept_images")), "export copy intact"
    _discard(cache_file)


@needs_fabio
def test_foreign_images_dir_preserved_and_never_served(tmp_path):
    """Export spares a user folder named like the sidecar, warns and never serves it as 2D."""
    zpath = _write_twod_zip(tmp_path)
    process_paths([zpath], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    cache_images = os.path.splitext(get_loaded_nxs_path())[0] + "_images"
    dest = str(tmp_path / "kept.nxs")
    foreign = tmp_path / "kept_images"
    os.makedirs(foreign)
    (foreign / "precious.txt").write_text("user data")
    (foreign / "image_001.edf").write_bytes(b"not an image")
    ok, msgs = export_nxs(dest)
    assert ok, "export itself must still succeed"
    assert any(m.startswith("Warning") for m in msgs), \
        "sidecar copy failure must be reported to the caller"
    assert (foreign / "precious.txt").read_text() == "user data", \
        "user's folder was clobbered by the sidecar copy"
    shutil.rmtree(cache_images)
    scans, _, _ = process_paths([dest], time_method=TimeMethod.ABSOLUTE,
                                data_source=DataSourceType.INHOUSE)
    assert make_twod_arrays(scans) == {}, \
        "foreign unmarked folder must not be served as 2D data"


@needs_fabio
def test_working_copy_carries_images_sidecar(tmp_path):
    """Add-echem to a reopened export keeps 2D viewable in the next export."""
    zpath = _write_twod_zip(tmp_path)
    process_paths([zpath], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    v1 = str(tmp_path / "master_v1.nxs")
    ok, _ = export_nxs(v1)
    assert ok
    process_paths([v1], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    extra = str(tmp_path / "extra_echem.txt")
    _builders.write_echem_txt(extra, n_rows=5)
    added = add_standard_echem_files([extra])
    assert added, "standard echem must append"
    v2 = str(tmp_path / "master_v2.nxs")
    ok, _ = export_nxs(v2)
    assert ok
    assert os.path.isdir(str(tmp_path / "master_v2_images")), \
        "sidecar lost through the working-copy flow"
    scans, _, _ = process_paths([v2], time_method=TimeMethod.ABSOLUTE,
                                data_source=DataSourceType.INHOUSE)
    arrays = make_twod_arrays(scans)
    assert 1 in arrays and arrays[1]["image"].shape == (8, 8), \
        "2D not viewable in the second-generation export"


@needs_fabio
def test_resolution_handles_foreign_separator_paths(tmp_path):
    """A recorded path with the other OS's separators still resolves."""
    zpath = _write_twod_zip(tmp_path)
    process_paths([zpath], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    dest = str(tmp_path / "kept.nxs")
    ok, _ = export_nxs(dest)
    assert ok
    process_paths([dest], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    recorded = "/tmp/operaxn_cache/cell_images/image_001.edf"  # POSIX form
    resolved = gui_input._resolve_twod_source(recorded)
    assert resolved and resolved.endswith("image_001.edf") and \
        os.path.isfile(resolved), f"got {resolved}"


# ============================================================================
# Session cache lifecycle
# ============================================================================

def test_cache_replaced_by_new_load(inhouse_dir, neutron_dir):
    """Loading a new dataset removes the previous unsaved cache file."""
    process_paths([inhouse_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    prev_cache = get_loaded_nxs_path()
    process_paths([neutron_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.NEUTRON)
    assert not os.path.exists(prev_cache)
    assert os.path.isfile(get_loaded_nxs_path())


def test_standard_echem_stored_and_survives_export(inhouse_dir, tmp_path):
    """Standard echem is stored in the .nxs, kept out of operando data and carried by the export."""
    std_echem_path = str(tmp_path / "standard_cycling.txt")
    _std_echem_file(std_echem_path)

    scans_se, echem_se, _ = process_paths(
        [inhouse_dir], time_method=TimeMethod.ABSOLUTE,
        data_source=DataSourceType.INHOUSE,
        standard_echem_files=[std_echem_path])
    model_s = core.load(get_loaded_nxs_path())
    assert len(model_s.standard_echem) == 1, \
        str([(e["source_file"], len(e["data"])) for e in model_s.standard_echem])
    assert len(model_s.standard_echem[0]["data"]) == 10
    assert model_s.standard_echem[0]["source_file"] == "standard_cycling.txt"
    assert abs(model_s.standard_echem[0]["data"]["echem_data"].iloc[0] - 4.0) < 1e-9

    assert len(echem_se) == 40, f"got {len(echem_se)}"
    assert scans_se[0]["echem"] == 3.76

    std_export = str(tmp_path / "with_std.nxs")
    export_nxs(std_export)
    assert len(core.load(std_export).standard_echem) == 1


def test_standard_echem_append_after_user_nxs(inhouse_dir, tmp_path):
    """Appending echem to a loaded user .nxs edits a cache copy, leaving the original untouched."""
    std_echem_path = str(tmp_path / "standard_cycling.txt")
    _std_echem_file(std_echem_path)
    process_paths([inhouse_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE,
                  standard_echem_files=[std_echem_path])
    std_export = str(tmp_path / "with_std.nxs")
    export_nxs(std_export)

    std_echem_later = str(tmp_path / "standard_added_later.txt")
    _std_echem_file(std_echem_later, n_rows=8, v0=3.5, i0=-0.5, date="06/03/2024")
    process_paths([std_export], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    added_later = add_standard_echem_files([std_echem_later])
    assert len(added_later) == 1
    assert os.path.abspath(get_loaded_nxs_path()) != os.path.abspath(std_export)
    assert len(core.load(std_export).standard_echem) == 1
    assert [item["source_file"]
            for item in core.load(get_loaded_nxs_path()).standard_echem] \
        == ["standard_cycling.txt", "standard_added_later.txt"]
    _discard(core.default_cache_path([inhouse_dir]))


@needs_fabio
def test_standard_echem_combines_with_embedded_twod(inhouse_edf_dir, tmp_path):
    """Standard echem and 2D inclusion can be combined at generation."""
    std_echem_path = str(tmp_path / "standard_cycling.txt")
    _std_echem_file(std_echem_path)
    process_paths([inhouse_edf_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE,
                  standard_echem_files=[std_echem_path],
                  include_2d_images=True, twod_max_display_size=0)
    m2 = core.load(get_loaded_nxs_path())
    assert len(m2.standard_echem) == 1
    assert m2.scans[0].twod_embedded and m2.scans[0].twod is not None


def test_upload_dialog_api():
    """The combined upload dialog exposes the expected headless API."""
    assert callable(UploadOptionsDialog)
    assert UploadOptionsDialog.DISPLAY_OPTIONS == \
        ["No downsampling", "4096", "2048", "1024", "512"]
    assert {"paths", "data_source", "time_method", "display_size", "include_2d",
            "title", "sample_name", "sample_description",
            "sample_preparation_date", "cycling_protocol",
            "standard_echem_files", "echem_time_tolerance"} \
        <= set(UploadOptions.__dataclass_fields__.keys())


def test_process_paths_threads_protocol_kwargs(inhouse_dir):
    """process_paths forwards preparation date and cycling protocol into the .nxs."""
    process_paths([inhouse_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE,
                  sample_preparation_date="2024-01-15",
                  cycling_protocol={"technique": "GCPL", "c_rate": "C/10"})
    meta = core.load(get_loaded_nxs_path()).global_metadata
    assert meta["cycling_protocol"]["technique"] == "GCPL", str(meta)
    assert meta["cycling_protocol"]["C_rate"] == "C/10", str(meta)
    assert str(meta["sample"]["preparation_date"]).startswith("2024-01-15"), \
        str(meta.get("sample"))


def test_process_paths_threads_tolerance(inhouse_dir):
    """The echem time tolerance reaches metadata and Excel rows; None gives the core default."""
    process_paths([inhouse_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE, echem_time_tolerance=120)
    meta = get_experiment_metadata()
    assert meta.get("echem_time_tolerance") == 120.0, str(meta)
    rows = experiment_metadata_rows(meta)
    assert {"field": "echem time tolerance", "value": "120 s"} in rows, rows
    process_paths([inhouse_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    assert get_experiment_metadata().get("echem_time_tolerance") == \
        core.ECHEM_TIME_TOLERANCE


@pytest.mark.display
def test_upload_dialog_tolerance_field(tk_root, monkeypatch, inhouse_dir, tmp_path):
    """The tolerance entry starts at the core default and accepts only finite positive numbers,
    as a float capped silently at core's limit; a directly opened .nxs greys it out and drops it."""
    errors = []
    monkeypatch.setattr("operaxn.dialog.messagebox.showerror",
                        lambda *args, **kwargs: errors.append(args))
    nxs = str(tmp_path / "direct.nxs")
    ok, msgs = core.generate([inhouse_dir], nxs, core.DataSourceType.INHOUSE)
    assert ok, str(msgs)

    dlg = UploadOptionsDialog(tk_root)
    try:
        assert dlg.tolerance_var.get() == f"{core.ECHEM_TIME_TOLERANCE:g}"
        dlg._paths = [inhouse_dir]
        dlg._update_option_states()
        assert str(dlg.tolerance_entry.cget("state")) == "normal"
        for bad in ("0", "-5", "abc", "nan", "inf"):
            dlg.tolerance_var.set(bad)
            dlg._confirm()
            assert dlg.result is None, bad
        assert len(errors) == 5 and all(a[0] == "Echem time tolerance" for a in errors)

        dlg._paths = [nxs]
        dlg._update_option_states()
        assert str(dlg.tolerance_entry.cget("state")) == "disabled"
        dlg.tolerance_var.set("abc")
        dlg._confirm()
        assert dlg.result is not None and dlg.result.echem_time_tolerance is None
        assert len(errors) == 5, "direct .nxs must not validate the greyed field"
    finally:
        if dlg.winfo_exists():
            dlg.destroy()

    dlg = UploadOptionsDialog(tk_root)
    try:
        dlg._paths = [inhouse_dir]
        dlg.tolerance_var.set(" 120.5 ")
        dlg._confirm()
        assert dlg.result is not None and dlg.result.echem_time_tolerance == 120.5
    finally:
        if dlg.winfo_exists():
            dlg.destroy()

    # Above core's limit the value quietly becomes the limit, with no error box
    from core.config import MAX_ECHEM_TIME_TOLERANCE as limit
    for text in ("1e12", f"{limit:g}"):
        dlg = UploadOptionsDialog(tk_root)
        try:
            dlg._paths = [inhouse_dir]
            dlg.tolerance_var.set(text)
            dlg._confirm()
            result = dlg.result
            assert result is not None and result.echem_time_tolerance == limit, text
        finally:
            if dlg.winfo_exists():
                dlg.destroy()
    assert len(errors) == 5, "the fallback is silent"

    dlg = UploadOptionsDialog(tk_root)
    try:
        dlg._paths = [inhouse_dir]
        dlg.tolerance_var.set("  ")
        dlg._confirm()
        assert dlg.result is not None and dlg.result.echem_time_tolerance is None, \
            "a blank field means the core default"
    finally:
        if dlg.winfo_exists():
            dlg.destroy()

    # Appending standard echem alone ignores the generation metadata, so it is not checked
    std = tmp_path / "std_echem.txt"
    _builders.write_echem_txt(str(std))
    dlg = UploadOptionsDialog(tk_root, has_loaded_data=True)
    try:
        dlg._std_echem = [str(std)]
        dlg.tolerance_var.set("abc")
        dlg.prep_date_var.set("notadate")
        dlg._confirm()
        assert dlg.result is not None and dlg.result.standard_echem_files == [str(std)]
        assert dlg.result.echem_time_tolerance is None and len(errors) == 5
    finally:
        if dlg.winfo_exists():
            dlg.destroy()


def test_display_size_caps_stored_synchrotron_images(tmp_path):
    """The single display-size setting also caps stored synchrotron images."""
    syn_dir = tmp_path / "synchrotron"
    syn_dir.mkdir()
    with h5py.File(str(syn_dir / "sample_123456.nxs"), "w") as f:
        f["/entry1/start_time"] = b"2024-02-05T10:05:00"
        f["/entry1/end_time"] = b"2024-02-05T10:07:00"
    with h5py.File(str(syn_dir / "sample_123456.hdf"), "w") as f:
        f["/entry/data"] = np.arange(256, dtype=np.float64).reshape(16, 16) + 1.0
    _builders.write_echem_txt(str(syn_dir / "echem.txt"))

    process_paths([str(syn_dir)], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.SYNCHROTRON,
                  include_2d_images=True, twod_max_display_size=8)
    m_syn = core.load(get_loaded_nxs_path())
    assert len(m_syn.scans) == 1, f"got {len(m_syn.scans)}"
    assert m_syn.scans[0].twod is not None, "no stored image"
    assert m_syn.scans[0].twod.shape == (8, 8), str(m_syn.scans[0].twod.shape)
    assert m_syn.scans[0].echem == 3.76, f"got {m_syn.scans[0].echem}"


def test_experiment_metadata_rows_bank_geometry(polaris_src):
    """Experiment rows list the flight path, each detector bank and each monitor element.
    A scalar detector field (lab EDF distance) or ragged monitor table yields no element rows."""
    assert experiment_metadata_rows({}) == []
    lab = {"instrument": {"name": "lab", "detector": {"distance": 0.1}}}
    assert [r["field"] for r in experiment_metadata_rows(lab)] == ["instrument"]
    synthetic = {"pre_sample_flightpath": 14.0,
                 "instrument": {"detector": {"detector_number": np.array([1, 2]),
                                             "distance": np.array([2.248, 1.783]),
                                             "polar_angle": np.array([10.4, 25.99])}},
                 "monitor": {"mode": "timer", "detector_number": np.array([671]),
                             "distance": np.array([3.27]),
                             "polar_angle": np.array([0.0])}}
    assert experiment_metadata_rows(synthetic) == [
        {"field": "pre_sample_flightpath", "value": "14.0"},
        {"field": "detector bank 1", "value": "L2 2.248 m, 2theta 10.40 deg"},
        {"field": "detector bank 2", "value": "L2 1.783 m, 2theta 25.99 deg"},
        {"field": "monitor mode", "value": "timer"},
        {"field": "monitor 671", "value": "3.270 m, 2theta 0.00 deg"}]
    # An in-house monitor (mode and preset, no element arrays) and a monitor
    # table whose arrays disagree in length yield no element rows
    inhouse = {"monitor": {"mode": "timer", "preset": 120.0}}
    assert experiment_metadata_rows(inhouse) == [
        {"field": "monitor mode", "value": "timer"},
        {"field": "monitor preset", "value": "120 s"}]
    ragged = {"monitor": {"mode": "timer", "detector_number": np.array([611, 612]),
                          "distance": np.array([7.641]),
                          "polar_angle": np.array([180.0, 180.0])}}
    assert experiment_metadata_rows(ragged) == [{"field": "monitor mode", "value": "timer"}]

    process_paths([polaris_src], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.NEUTRON)
    fields = {r["field"]: r["value"]
              for r in experiment_metadata_rows(get_experiment_metadata())}
    assert fields["instrument"] == "POLARIS", str(fields)
    assert fields["pre_sample_flightpath"] == "14.0", str(fields)
    assert fields["detector bank 5"] == "L2 1.251 m, 2theta 146.72 deg", str(fields)
    assert fields["monitor mode"] == "timer", str(fields)
    assert fields["monitor 611"] == "7.641 m, 2theta 180.00 deg", str(fields)


def test_cli_profile_option(tmp_path, capsys):
    """operaxn --profile prints a Mantid file's INSTRUMENT_PROFILES block, exits 1 if unreadable."""
    path = str(tmp_path / "TESTINST1-2.nxs")
    _builders.write_mantid_processed_file(path)
    assert cli_main(["--profile", path]) == 0
    out = capsys.readouterr().out
    assert '"testinst": {' in out and '"pre_sample_flightpath_m": 14.0' in out
    assert '"monitor_mode": None' in out and '"monitors": {' in out
    assert cli_main(["--profile", str(tmp_path / "missing.nxs")]) == 1
    assert "error:" in capsys.readouterr().err


def test_excel_sheet_widths_survive_nan_and_empty_frames():
    """Column fitting handles NaN (pandas 3 keeps it through astype(str)) and a header-only sheet."""
    from openpyxl import Workbook
    frames = (pd.DataFrame({"Voltage (V)": pd.Series([3.71, np.nan], dtype=object),
                            "note": ["a", None]}),
              pd.DataFrame(columns=["scan number", "Has Echem"]))
    for df in frames:
        worksheet = Workbook().active
        worksheet.append(list(df.columns))
        OPERAXN._format_sheet(worksheet, df)
        widths = [worksheet.column_dimensions[c].width for c in ("A", "B")]
        assert widths == [len(col) + 3 for col in df.columns], widths
        assert worksheet["A1"].font.bold and worksheet.freeze_panes == "A2"


def test_neutron_statistics_sheet_reads_echem_columns(tmp_path):
    """Neutron Statistics rows carry the correlated voltage and current (mA); uncorrelated say No."""
    def scan(num, volts, amps):
        return {"scan_num": num, "timestamp": "2026-06-01 10:00:00", "echem": volts,
                "current": amps, "echem_timestamp": "", "exposure_time": 120.0,
                "neutron_start": "2026-06-01 10:00:00", "neutron_end": "2026-06-01 10:02:00",
                "neutron_files": {"1": {"tof": f"/data/run{num}_b1_tof.dat"}}, "neutron_meta": None}
    fake = SimpleNamespace(state=SimpleNamespace(
        scans=[scan(1, 3.71, 0.25), scan(2, None, None)],
        data_source=core.DataSourceType.NEUTRON, time_method="absolute"))
    summary = pd.DataFrame([OPERAXN._create_scan_summary(fake, s) for s in fake.state.scans])
    rows = [OPERAXN._calculate_neutron_statistics(fake, r) for _, r in summary.iterrows()]
    assert rows[0]["Voltage (V)"] == 3.71 and rows[0]["Current (mA)"] == 0.25
    assert rows[0]["Has Echem"] == "Yes" and rows[0]["Scan Duration"] == "120.0 s"
    assert rows[1]["Voltage (V)"] == "" and rows[1]["Current (mA)"] == ""
    assert rows[1]["Has Echem"] == "No"
    assert "Current (A)" not in rows[0], "values are mA"

    # The sheet gets the shared formatting: bold frozen header, capped widths
    from openpyxl import load_workbook
    fake._calculate_neutron_statistics = (
        lambda row: OPERAXN._calculate_neutron_statistics(fake, row))
    fake._format_sheet = OPERAXN._format_sheet
    book = tmp_path / "stats.xlsx"
    with pd.ExcelWriter(book, engine="openpyxl") as writer:
        OPERAXN._add_neutron_statistics_sheet(fake, writer, summary)
    sheet = load_workbook(book)["Neutron Statistics"]
    assert sheet["A1"].font.bold and sheet.freeze_panes == "A2"
    assert all(dim.width <= 50 for dim in sheet.column_dimensions.values())


def test_excel_export_asks_for_the_file_before_its_thread(monkeypatch):
    """The save dialog runs on the calling (Tk) thread; cancelling it starts no export thread."""
    import threading
    from tkinter import filedialog
    import operaxn.gui as gui_module
    asked, started = [], []
    monkeypatch.setattr(filedialog, "asksaveasfilename",
                        lambda **_k: asked.append(threading.current_thread()) or "")
    monkeypatch.setattr(gui_module.threading, "Thread",
                        lambda *a, **k: started.append(k) or SimpleNamespace(start=lambda: None))
    fake = SimpleNamespace(master=None, state=SimpleNamespace(
        scans=[{"scan_num": 1}], data_source=core.DataSourceType.INHOUSE),
        _warn_if_busy=lambda action: False)
    OPERAXN._export_to_excel(fake)
    assert asked == [threading.current_thread()] and started == []


def test_config_summary_logs_after_setup(caplog):
    """The debug configuration summary logs its lines when called after logging setup."""
    from operaxn.config import log_config_summary
    with caplog.at_level(logging.DEBUG, logger="operaxn.config"):
        log_config_summary()
    text = "\n".join(r.getMessage() for r in caplog.records)
    assert "OperaXN Configuration" in text and "Max Workers:" in text


def test_splash_fill_stops_at_track_end():
    """The splash progress fill is drawn no further than 100 % of its track."""
    from operaxn.main import EnhancedSplashScreen

    class Var:
        def __init__(self, value):
            self.value = value

        def get(self):
            return self.value

        def set(self, value):
            self.value = value

    drawn = []

    class Canvas:
        def coords(self, item, *xy):
            if xy:
                drawn.append(xy)
            return [50, 10, 60, 20]

        def itemconfig(self, *_args, **_kwargs):
            pass

    fake = SimpleNamespace(
        animation_running=True, splash=SimpleNamespace(winfo_exists=lambda: True),
        progress_var=Var(99), canvas=Canvas(), _schedule_animation=lambda *_args: None,
        _animate_progress=None,  # rescheduling target; the stub never calls it
        progress_target=100,
        ui_elements={"bar_width": 300, "bar_center_x": 200,
                     "progress_rect": 1, "progress_glow": 2})
    EnhancedSplashScreen._animate_progress(fake)
    assert fake.progress_var.get() == 100
    assert drawn[0][2] == 350, "the fill must end at the track's right edge (50 + 300)"


@pytest.mark.display
def test_plot_settings_apply_redraws_once(tk_root):
    """Applying plot settings redraws through the update callback exactly once."""
    from operaxn.dialog import PlotSettingsDialog
    from operaxn.gui import VisualiserConfig
    calls = []
    dlg = PlotSettingsDialog(tk_root, VisualiserConfig(), lambda: calls.append(1))
    try:
        dlg._apply_settings()
    finally:
        dlg.destroy()
    assert calls == [1]


def test_direct_nxs_opens_as_its_recorded_source(neutron_dir, tmp_path):
    """A directly opened .nxs is read and displayed as the source it records,
    whatever was selected; an unreadable path gives None."""
    process_paths([neutron_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.NEUTRON)
    neutron_nxs = str(tmp_path / "neutron_run.nxs")
    shutil.copy(get_loaded_nxs_path(), neutron_nxs)
    assert core.read_data_source(neutron_nxs) == "neutron"
    assert core.read_data_source(str(tmp_path / "missing.nxs")) is None

    scans, _, _ = process_paths([neutron_nxs], time_method=TimeMethod.ABSOLUTE,
                                data_source=DataSourceType.INHOUSE)
    assert get_loaded_data_source() == DataSourceType.NEUTRON
    assert make_neutron_arrays(scans), "neutron banks must load from a directly opened file"


def test_working_copy_keeps_the_export_stem(inhouse_dir, tmp_path):
    """Adding echem to an opened .nxs edits <stem>_working.nxs, but exports keep <stem>."""
    from operaxn.gui import _experiment_stem
    process_paths([inhouse_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    user_nxs = str(tmp_path / "MyCell_run3.nxs")
    export_nxs(user_nxs)
    process_paths([user_nxs], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE)
    assert get_experiment_stem() == "MyCell_run3"

    std_echem_path = str(tmp_path / "standard_cycling.txt")
    _std_echem_file(std_echem_path)
    assert add_standard_echem_files([std_echem_path])
    assert os.path.basename(get_loaded_nxs_path()) == "MyCell_run3_working.nxs"
    assert get_experiment_stem() == "MyCell_run3" and _experiment_stem() == "MyCell_run3"
    _discard(core.default_cache_path([inhouse_dir]))


def test_reset_session_forgets_the_experiment(inhouse_dir, tmp_path):
    """reset_session drops the unsaved cache file and every session field, so an echem
    append afterwards has no experiment to target."""
    std_echem_path = str(tmp_path / "standard_cycling.txt")
    _std_echem_file(std_echem_path)
    process_paths([inhouse_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.INHOUSE,
                  standard_echem_files=[std_echem_path])
    cache_file = get_loaded_nxs_path()
    assert get_standard_echem()
    reset_session()
    assert not os.path.exists(cache_file)
    assert get_loaded_nxs_path() is None and get_experiment_stem() is None
    assert get_standard_echem() == [] and get_experiment_metadata() == {}
    assert get_loaded_data_source() is None
    with pytest.raises(RuntimeError):
        add_standard_echem_files([std_echem_path])


def _gui_app(tk_root, monkeypatch):
    """Return (window, app, messages), message boxes recorded. Workers run on the Tk thread:
    after() from another thread needs a running mainloop, and tests pump with update()."""
    from tkinter import Toplevel
    import operaxn.gui as gui_module
    win = Toplevel(tk_root)

    class DeferredThread:
        def __init__(self, target, daemon=None):
            self._target = target

        def start(self):
            win.after(0, self._target)

    monkeypatch.setattr(gui_module, "threading", SimpleNamespace(Thread=DeferredThread))
    app = OPERAXN(master=win)
    messages = []
    monkeypatch.setattr(app, "_show_message", lambda *args: messages.append(args))
    return win, app, messages


def _wait_for_worker(win, app, timeout=120.0):
    """Pump Tk events until the app's worker task (load, append or export) has finished."""
    import time
    deadline = time.monotonic() + timeout
    while app._is_busy():
        assert time.monotonic() < deadline, "worker task did not finish"
        win.update()
        time.sleep(0.01)
    win.update()


@pytest.mark.display
def test_upload_and_clear_locked_while_loading(tk_root, monkeypatch, neutron_dir, tmp_path):
    """Upload and Clear are disabled and refused while a load runs, and come back after;
    a neutron .nxs opened under the default Laboratory source still loads as neutron."""
    process_paths([neutron_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.NEUTRON)
    neutron_nxs = str(tmp_path / "neutron_run.nxs")
    shutil.copy(get_loaded_nxs_path(), neutron_nxs)

    win, app, messages = _gui_app(tk_root, monkeypatch)
    try:
        opened = []
        import operaxn.gui as gui_module
        monkeypatch.setattr(gui_module, "UploadOptionsDialog",
                            lambda *_a, **_k: opened.append(1))
        options = UploadOptions(paths=[neutron_nxs], data_source=DataSourceType.INHOUSE,
                                time_method=TimeMethod.ABSOLUTE, display_size=0,
                                include_2d=False)
        app._process_files_async(options)
        buttons = app.button_panel.buttons
        assert app._is_busy()
        assert str(buttons["upload_files"].cget("state")) == "disabled"
        assert str(buttons["clear_all"].cget("state")) == "disabled"
        app._upload_files()
        assert opened == [], "Upload must not open a dialog mid-load"
        assert app._clear_all() is False, "Clear must be refused mid-load"

        _wait_for_worker(win, app)
        assert str(buttons["upload_files"].cget("state")) == "normal"
        assert str(buttons["clear_all"].cget("state")) == "normal"
        assert app.state.data_source == DataSourceType.NEUTRON
        assert app.state.neutron_arrays, messages
        assert messages and messages[-1][0] == "Success", messages
    finally:
        if win.winfo_exists():
            win.destroy()


@pytest.mark.display
def test_replot_closes_old_figure_and_rebinds_plot_settings(tk_root, monkeypatch, inhouse_dir):
    """Each Plot frees the figure it replaces, an open Plot Settings dialog reopens on the
    new config, and Clear frees the figure and forgets the session."""
    import matplotlib.pyplot as plt
    win, app, messages = _gui_app(tk_root, monkeypatch)
    try:
        options = UploadOptions(paths=[inhouse_dir], data_source=DataSourceType.INHOUSE,
                                time_method=TimeMethod.ABSOLUTE, display_size=0,
                                include_2d=False)
        app._process_files_async(options)
        _wait_for_worker(win, app)
        assert app.state.scans, messages

        def open_figures():
            # Compare objects: pyplot hands a closed figure's number to the next one
            return [manager.canvas.figure
                    for manager in plt._pylab_helpers.Gcf.get_all_fig_managers()]

        before = len(open_figures())
        app._plot_data()
        first = app.fig
        app._show_plot_settings()
        old_dialog = app._plot_settings_dialog
        app._plot_data()
        assert app.fig is not first, messages
        assert all(fig is not first for fig in open_figures()), "replaced figure not closed"
        assert len(open_figures()) == before + 1
        new_dialog = app._plot_settings_dialog
        assert new_dialog is not None and new_dialog is not old_dialog
        assert not old_dialog.winfo_exists()
        assert new_dialog.config is app.config, "Plot Settings must edit the live config"

        monkeypatch.setattr("operaxn.gui.messagebox.askyesno", lambda *_a, **_k: True)
        last = app.fig
        assert app._clear_all()
        assert all(fig is not last for fig in open_figures()), "Clear must close the figure"
        assert get_loaded_nxs_path() is None and get_standard_echem() == []
    finally:
        if win.winfo_exists():
            win.destroy()


def _load_inhouse(win, app, inhouse_dir):
    """Run a raw in-house load through the app's worker and wait for it."""
    app._process_files_async(UploadOptions(paths=[inhouse_dir],
                                           data_source=DataSourceType.INHOUSE,
                                           time_method=TimeMethod.ABSOLUTE,
                                           display_size=0, include_2d=False))
    _wait_for_worker(win, app)
    assert app.state.scans


@pytest.mark.display
def test_failed_load_unlocks_upload_and_clear(tk_root, monkeypatch, inhouse_dir):
    """A load whose processing raises shows the error and leaves Upload and Clear usable."""
    import operaxn.gui as gui_module
    win, app, messages = _gui_app(tk_root, monkeypatch)
    try:
        def boom(*_a, **_k):
            raise RuntimeError("unreadable input")

        monkeypatch.setattr(gui_module, "process_paths", boom)
        app._process_files_async(UploadOptions(paths=[inhouse_dir],
                                               data_source=DataSourceType.INHOUSE,
                                               time_method=TimeMethod.ABSOLUTE,
                                               display_size=0, include_2d=False))
        _wait_for_worker(win, app)
        assert messages and messages[-1][0] == "Error" and "unreadable input" in messages[-1][1]
        buttons = app.button_panel.buttons
        assert str(buttons["upload_files"].cget("state")) == "normal"
        assert str(buttons["clear_all"].cget("state")) == "normal"
    finally:
        if win.winfo_exists():
            win.destroy()


@pytest.mark.display
def test_plot_keeps_or_ends_a_worker_busy_state_correctly(tk_root, monkeypatch, inhouse_dir):
    """Plot during an echem append keeps the append's lock; if the append finishes inside
    Plot's error box, the lock ends rather than coming back when the box closes."""
    from operaxn.gui import UIState
    win, app, messages = _gui_app(tk_root, monkeypatch)
    try:
        _load_inhouse(win, app, inhouse_dir)

        app._set_busy(UIState.LOADING)  # an append is running
        app._plot_data()
        assert app._is_busy(), "a successful Plot must not end the append's lock"

        # The append's completion runs while Plot's error box pumps events
        def failing_figure(*_a, **_k):
            raise RuntimeError("draw failed")

        def message_box(*args):
            messages.append(args)
            app._finish_busy(None)

        monkeypatch.setattr(app, "_create_figure", failing_figure)
        monkeypatch.setattr(app, "_show_message", message_box)
        app._plot_data()
        assert messages[-1][0] == "Error"
        assert not app._is_busy(), "the finished append must stay finished"
        assert str(app.button_panel.buttons["clear_all"].cget("state")) == "normal"
    finally:
        if win.winfo_exists():
            win.destroy()


@pytest.mark.display
def test_echem_append_while_busy_is_refused_with_a_message(tk_root, monkeypatch, tmp_path):
    """Echem picked while an export started (e.g. during ICI's file picker) is not dropped
    silently: the user is told to add it again."""
    from operaxn.gui import UIState
    import operaxn.gui as gui_module
    win, app, messages = _gui_app(tk_root, monkeypatch)
    try:
        started = []
        monkeypatch.setattr(gui_module, "add_standard_echem_files", started.append)
        app._set_busy(UIState.EXPORTING)
        app._append_standard_echem_async([str(tmp_path / "std.txt")])
        win.update()
        assert started == [] and messages[-1][0] == "Busy", messages
        assert "add electrochemistry" in messages[-1][1]
    finally:
        if win.winfo_exists():
            win.destroy()


@pytest.mark.display
def test_echem_only_upload_needs_a_loaded_experiment(tk_root, monkeypatch, tmp_path):
    """Standard echem alone (dialog opened with data that a Clear then removed) is refused."""
    import operaxn.gui as gui_module
    win, app, messages = _gui_app(tk_root, monkeypatch)
    try:
        options = UploadOptions(paths=[], data_source=DataSourceType.INHOUSE,
                                time_method=TimeMethod.ABSOLUTE, display_size=0,
                                include_2d=False,
                                standard_echem_files=[str(tmp_path / "std.txt")])
        monkeypatch.setattr(gui_module, "UploadOptionsDialog",
                            lambda *_a, **_k: SimpleNamespace(get_result=lambda: options))
        appended = []
        monkeypatch.setattr(app, "_append_standard_echem_async", appended.append)
        app._upload_files()
        assert appended == [] and messages[-1][0] == "No experiment loaded", messages
    finally:
        if win.winfo_exists():
            win.destroy()


@pytest.mark.display
def test_exports_lock_upload_and_clear(tk_root, monkeypatch, inhouse_dir, tmp_path):
    """NeXus and Excel exports hold the busy lock, so Clear cannot delete the file being
    copied or the scans being summarised."""
    from operaxn.gui import UIState
    win, app, messages = _gui_app(tk_root, monkeypatch)
    try:
        _load_inhouse(win, app, inhouse_dir)
        cache_file = get_loaded_nxs_path()
        target = {"path": str(tmp_path / "exported.nxs")}
        monkeypatch.setattr("operaxn.gui.filedialog.asksaveasfilename",
                            lambda *_a, **_k: target["path"])
        monkeypatch.setattr("operaxn.gui.messagebox.askyesno", lambda *_a, **_k: True)

        app._export_nxs()
        assert app.state.ui_state == UIState.EXPORTING
        assert app._clear_all() is False and os.path.isfile(cache_file)
        app._export_nxs()
        assert messages[-1][0] == "Busy", "a second export must wait"
        _wait_for_worker(win, app)
        assert messages[-1][0] == "Success", messages
        assert os.path.isfile(target["path"])

        target["path"] = str(tmp_path / "summary.csv")
        app._export_to_excel()
        assert app.state.ui_state == UIState.EXPORTING
        assert app._clear_all() is False and app.state.scans
        _wait_for_worker(win, app)
        assert messages[-1][0] == "Success", messages
        assert os.path.isfile(target["path"])

        assert str(app.button_panel.buttons["clear_all"].cget("state")) == "normal"
        assert app._clear_all() and os.path.isfile(str(tmp_path / "exported.nxs"))
        _discard(cache_file)
    finally:
        if win.winfo_exists():
            win.destroy()


@pytest.mark.display
def test_upload_dialog_takes_source_from_nxs(tk_root, neutron_dir, inhouse_dir, tmp_path):
    """Selecting a directly opened .nxs shows and locks its recorded source, with a grey note;
    choosing raw data again restores the user's own choice."""
    from operaxn.config import OPERAXNTheme
    process_paths([neutron_dir], time_method=TimeMethod.ABSOLUTE,
                  data_source=DataSourceType.NEUTRON)
    neutron_nxs = str(tmp_path / "neutron_run.nxs")
    shutil.copy(get_loaded_nxs_path(), neutron_nxs)

    dlg = UploadOptionsDialog(tk_root)
    try:
        dlg.source_var.set("synchrotron")
        dlg._paths = [neutron_nxs]
        dlg._update_option_states()
        assert dlg.source_var.get() == "neutron"
        assert all(str(r.cget("state")) == "disabled" for r in dlg._source_radios)
        assert "source and generation options" in dlg.note_var.get()
        assert dlg._note_label.cget("fg") == OPERAXNTheme.COLORS['text_secondary']

        dlg._paths = [inhouse_dir]
        dlg._update_option_states()
        assert dlg.source_var.get() == "synchrotron"
        assert all(str(r.cget("state")) == "normal" for r in dlg._source_radios)
        assert dlg.note_var.get() == ""
    finally:
        if dlg.winfo_exists():
            dlg.destroy()
